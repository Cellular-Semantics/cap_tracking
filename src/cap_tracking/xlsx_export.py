"""Generate a formatted .xlsx review copy of the tracker CSV.

The CSV stays the source of truth — this is a read-only, regeneratable
export for reviewing/downloading in Excel: a real Excel Table (filter/sort
dropdowns, banded rows), bold frozen header, columns sized to fit their
content in full, and each row's font colored by completion status.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from cap_tracking.schema import FIELDNAMES

_FONT_COLORS = {
    "green": "FF008000",
    "amber": "FFC08000",
    "red": "FFFF0000",
    "grey": "FF808080",
}


def row_color_category(row: dict) -> str:
    """One of 'green' (completed), 'amber' (in progress), 'red' (untouched),
    'grey' (skip/blocked — not part of the completion lifecycle)."""
    route = row.get("action_route")

    if route in ("skip", "blocked"):
        return "grey"

    if route == "cap_comment":
        if row.get("comment_status") == "done":
            return "green"
        if row.get("comment_status") == "drafted":
            return "amber"
        return "red"

    if route == "github_issue":
        if row.get("github_issue_status") == "filed" and row.get("github_ticket"):
            return "green"
        if row.get("github_issue_status") == "drafted":
            return "amber"
        return "red"

    if route == "both":
        gh_done = row.get("github_issue_status") == "filed" and bool(row.get("github_ticket"))
        cap_done = row.get("comment_status") == "done"
        if gh_done and cap_done:
            return "green"
        gh_started = row.get("github_issue_status") in ("drafted", "filed")
        cap_started = row.get("comment_status") in ("drafted", "done")
        if gh_started or cap_started:
            return "amber"
        return "red"

    return "grey"


def export_xlsx(rows: list[dict], out_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "findings"

    ws.append(FIELDNAMES)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"

    widths = [len(name) for name in FIELDNAMES]
    for row in rows:
        values = [row.get(field, "") for field in FIELDNAMES]
        ws.append(values)
        color = row_color_category(row)
        font = Font(color=_FONT_COLORS[color])
        for cell in ws[ws.max_row]:
            cell.font = font
        for i, value in enumerate(values):
            widths[i] = max(widths[i], len(str(value)))

    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width + 2

    last_col = get_column_letter(len(FIELDNAMES))
    last_row = len(rows) + 1
    table = Table(displayName="FindingsTable", ref=f"A1:{last_col}{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9", showRowStripes=True, showFirstColumn=False
    )
    ws.add_table(table)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
